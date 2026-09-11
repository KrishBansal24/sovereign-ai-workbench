"""Integration tests for validated local document upload and extraction routes."""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from httpx import Response
from pypdf import PdfWriter
from pytest import MonkeyPatch

from app.api.routes import documents as documents_route
from app.main import app
from app.services.documents.document_service import DocumentService


client = TestClient(app)


def use_temporary_store(
    monkeypatch: MonkeyPatch,
    tmp_path: Path,
    max_upload_size_mb: int = 20,
) -> DocumentService:
    """Replace the route's production store with isolated test storage."""
    service = DocumentService(data_directory=tmp_path, max_upload_size_mb=max_upload_size_mb)
    monkeypatch.setattr(documents_route, "document_service", service)
    return service


def upload(filename: str, content: bytes, content_type: str) -> Response:
    """Submit one multipart document upload through the real API route."""
    return client.post(
        "/api/documents/upload",
        files={"file": (filename, content, content_type)},
    )


def test_upload_txt_and_retrieve_text(monkeypatch, tmp_path):
    use_temporary_store(monkeypatch, tmp_path)

    response = upload("inspection.txt", b"Pump inspection complete.\r\nNo leaks found.", "text/plain")

    assert response.status_code == 201
    metadata = response.json()
    assert metadata["file_type"] == "txt"
    assert metadata["extraction_status"] == "extracted"
    text_response = client.get(f"/api/documents/{metadata['document_id']}/text")
    assert text_response.status_code == 200
    assert "No leaks found." in text_response.json()["text"]


def test_upload_valid_pdf_reports_ocr_requirement_when_no_text(monkeypatch, tmp_path):
    use_temporary_store(monkeypatch, tmp_path)
    # A blank valid PDF exercises the scanned/text-poor extraction status.
    pdf = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(pdf)

    response = upload("synthetic.pdf", pdf.getvalue(), "application/pdf")

    assert response.status_code == 201
    assert response.json()["file_type"] == "pdf"
    assert response.json()["page_count"] == 1
    assert response.json()["extraction_status"] == "ocr_required"


def test_upload_docx_extracts_paragraphs_and_tables(monkeypatch, tmp_path):
    use_temporary_store(monkeypatch, tmp_path)
    document = Document()
    document.add_paragraph("Synthetic maintenance report")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Equipment"
    table.cell(0, 1).text = "Pump A"
    buffer = BytesIO()
    document.save(buffer)

    response = upload(
        "report.docx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 201
    document_id = response.json()["document_id"]
    text = client.get(f"/api/documents/{document_id}/text").json()["text"]
    assert "Synthetic maintenance report" in text
    assert "Equipment\tPump A" in text


def test_upload_rejects_unsupported_empty_and_oversized_files(monkeypatch, tmp_path):
    use_temporary_store(monkeypatch, tmp_path, max_upload_size_mb=1)

    assert upload("drawing.png", b"not an image", "image/png").status_code == 422
    assert upload("empty.txt", b"", "text/plain").status_code == 422
    assert upload("large.txt", b"x" * (1024 * 1024 + 1), "text/plain").status_code == 422


def test_safe_storage_list_and_delete(monkeypatch, tmp_path):
    service = use_temporary_store(monkeypatch, tmp_path)

    response = upload("../../unsafe-name.txt", b"synthetic test document", "text/plain")

    assert response.status_code == 201
    metadata = response.json()
    assert "/" not in metadata["filename"]
    assert "\\" not in metadata["filename"]
    assert (service.files_directory / f"{metadata['document_id']}.txt").exists()
    assert client.get("/api/documents").json()["documents"][0]["document_id"] == metadata["document_id"]
    assert client.delete(f"/api/documents/{metadata['document_id']}").status_code == 200
    assert client.get(f"/api/documents/{metadata['document_id']}").status_code == 404


def test_missing_and_corrupt_documents_fail_safely(monkeypatch, tmp_path):
    use_temporary_store(monkeypatch, tmp_path)

    assert client.get("/api/documents/not-a-document-id").status_code == 404
    assert upload("corrupt.pdf", b"%PDF-not-a-real-pdf", "application/pdf").status_code == 422
    assert upload(
        "corrupt.docx",
        b"PKnot-a-docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ).status_code == 422
