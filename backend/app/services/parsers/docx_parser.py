"""Deterministic local DOCX extraction used only by DocumentService."""

from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.services.parsers.text_parser import normalize_text


class DocxExtractionError(Exception):
    """Raised when an uploaded DOCX package cannot be opened safely."""


def extract_text(content: bytes) -> str:
    """Extract paragraphs and table rows from validated DOCX bytes.

    Args:
        content: Raw bytes already validated by DocumentService.

    Returns:
        Conservatively normalized plain text.
    """
    try:
        document = Document(BytesIO(content))
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as error:
        raise DocxExtractionError("The DOCX file is malformed or unreadable.") from error

    # Tables are preserved as tab-separated rows rather than discarded.
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return normalize_text("\n\n".join(parts))
